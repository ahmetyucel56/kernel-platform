"""Sinif not dosyasi (Excel) — akademisyen notlari OBS'ye aktarmak icin indirir.

- "Notlar" sayfasi: Ogrenci No, Ad Soyad, odev basina not, notlanan odevlerin
  ortalamasi. Not verilmemis hucre BOS kalir (0 yazilmaz; karar hocanindir).
- "Ayrinti" sayfasi: teslim durumu (zamaninda / uzatmada / teslim yok), surum,
  tarih, not verilen surum, AI kapsam, Clean Code, intihal benzerligi, uyarilar.

Indirme, proje ZIP'indeki gibi 5 dk gecerli imzali link ile (purpose=gradebook).
"""
from __future__ import annotations

import io
import re
import uuid
from datetime import datetime, timedelta, timezone

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.db import get_db
from app.deps import require_roles
from app.lib.classinfo import course_names
from app.models import AiAnalysis, Assignment, Class, Enrollment, Score, Submission, User
from app.routers.insights import _latest_req_analyses, _owned_class
from app.routers.submissions import _aware, _effective_deadline, _slug
from app.security import DOWNLOAD_TTL_MINUTES, create_download_token, decode_token
from app.services import analysis_service as ana

router = APIRouter(tags=["gradebook"])

TR = timezone(timedelta(hours=3))  # Turkiye (2016'dan beri sabit UTC+3)
HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1B191F")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def _local(dt: datetime | None) -> datetime | None:
    # Excel saat dilimi tutmaz: Turkiye saatine cevirip "naive" yaz.
    return _aware(dt).astimezone(TR).replace(tzinfo=None) if dt else None


def _latest_by_type(db: Session, sub_id: uuid.UUID, kind: str) -> AiAnalysis | None:
    return db.scalar(
        select(AiAnalysis)
        .where(AiAnalysis.target_submission_id == sub_id, AiAnalysis.analysis_type == kind)
        .order_by(AiAnalysis.created_at.desc())
        .limit(1)
    )


def _rows(db: Session, cls: Class):
    """Ogrenci x odev hucreleri: not + ayrinti."""
    students = db.scalars(
        select(User).join(Enrollment, Enrollment.student_id == User.id)
        .where(Enrollment.class_id == cls.id).order_by(User.school_no, User.full_name)
    ).all()
    assignments = db.scalars(
        select(Assignment).where(Assignment.class_id == cls.id).order_by(Assignment.deadline_at)
    ).all()
    cells: dict[tuple, dict] = {}
    for a in assignments:
        subs = db.scalars(select(Submission).where(Submission.assignment_id == a.id)).all()
        by_student: dict[uuid.UUID, list[Submission]] = {}
        for s in subs:
            by_student.setdefault(s.student_id, []).append(s)
        latest = {sid: max(v, key=lambda x: x.version_number) for sid, v in by_student.items()}
        req = _latest_req_analyses(db, [s.id for s in latest.values()])
        for st in students:
            versions = by_student.get(st.id, [])
            sub = latest.get(st.id)
            cell = {"assignment": a, "sub": sub, "score": None, "graded_version": None,
                    "status": "Teslim yok", "coverage": None, "clean": None, "similarity": None,
                    "warnings": []}
            if sub:
                base = _aware(a.deadline_at)
                at = _aware(sub.submitted_at)
                if at <= base:
                    cell["status"] = "Zamanında"
                elif at <= _effective_deadline(db, a, st.id):
                    cell["status"] = "Uzatma süresinde"
                else:  # pragma: no cover - yukleme ucu buna izin vermez
                    cell["status"] = "Geç"
                # Hocanin verdigi EN SON not (hangi surume verilmis olursa olsun)
                sc = db.scalar(
                    select(Score).where(Score.submission_id.in_([v.id for v in versions]))
                    .order_by(Score.graded_at.desc()).limit(1)
                )
                if sc:
                    cell["score"] = sc.score
                    graded = next(v for v in versions if v.id == sc.submission_id)
                    cell["graded_version"] = graded.version_number
                    if graded.version_number < sub.version_number:
                        cell["warnings"].append(
                            f"Notlandıktan sonra yeni sürüm yüklendi (not v{graded.version_number}, son v{sub.version_number})")
                an = req.get(sub.id)
                if an is not None:
                    if ana.requirements_changed(an.detail_json, list(a.requirements_json or [])):
                        cell["warnings"].append("Gereksinim analizi eski kurallarla yapılmış")
                    else:
                        cell["coverage"] = (an.summary_json or {}).get("coverage")
                cc = _latest_by_type(db, sub.id, ana.CLEAN_CODE)
                if cc:
                    cell["clean"] = (cc.summary_json or {}).get("score")
                pl = _latest_by_type(db, sub.id, ana.PLAGIARISM)
                if pl:
                    sim = (pl.summary_json or {}).get("top_similarity")
                    cell["similarity"] = sim
                    if isinstance(sim, (int, float)) and sim >= ana.SIMILARITY_WARN:
                        cell["warnings"].append(f"Yüksek benzerlik: %{sim}")
            cells[(st.id, a.id)] = cell
    return students, assignments, cells


def _sheet_title(name: str, used: set[str]) -> str:
    """Excel sayfa adı: en fazla 31 karakter, []:*?/\\ yok, tekrar yok."""
    base = re.sub(r"[\[\]:*?/\\]", " ", name).strip()[:31] or "Notlar"
    title, n = base, 2
    while title in used:
        suffix = f" ({n})"
        title, n = base[: 31 - len(suffix)] + suffix, n + 1
    used.add(title)
    return title


def build_gradebook(db: Session, cls: Class) -> bytes:
    students, assignments, cells = _rows(db, cls)
    cnames = course_names(db, [a.course_id for a in assignments])
    wb = Workbook()

    # --- Notlar: her ders ayrı sayfa (ortalama dersin kendi ödevlerinden) ---
    groups: dict[str, list] = {}
    for a in assignments:
        groups.setdefault(cnames.get(a.course_id) or "Diğer", []).append(a)
    if not groups:
        groups["Notlar"] = []
    used: set[str] = set()
    for i, (course, items) in enumerate(sorted(groups.items())):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = _sheet_title("Notlar" if len(groups) == 1 else course, used)
        head = ["Öğrenci No", "Ad Soyad"] + [a.title for a in items] + ["Ortalama*"]
        ws.append(head)
        for st in students:
            scores = [cells[(st.id, a.id)]["score"] for a in items]
            given = [s for s in scores if s is not None]
            avg = round(sum(given) / len(given), 1) if given else None
            ws.append([st.school_no or "", st.full_name] + scores + [avg])
        ws.append([])
        ws.append(["* Ortalama yalnızca bu dersin not verilmiş ödevlerinden hesaplanır. Boş hücre: not "
                   "verilmemiş (teslim edilmemiş ya da henüz notlanmamış)."])
        ws.append([f"{cls.name} · {course} · Oluşturulma: {datetime.now(TR):%d.%m.%Y %H:%M}"])
        _style(ws, len(head), widths=[14, 24] + [18] * len(items) + [12])

    # --- Ayrinti ---
    wd = wb.create_sheet("Ayrıntı")
    dh = ["Öğrenci No", "Ad Soyad", "Ders", "Ödev", "Son teslim tarihi", "Durum", "Son sürüm",
          "Teslim zamanı", "Not", "Not verilen sürüm", "AI kapsam (%)", "Clean Code",
          "En yüksek benzerlik (%)", "Uyarı"]
    wd.append(dh)
    for st in students:
        for a in assignments:
            c = cells[(st.id, a.id)]
            sub = c["sub"]
            wd.append([st.school_no or "", st.full_name, cnames.get(a.course_id) or "", a.title,
                       _local(a.deadline_at), c["status"],
                       sub.version_number if sub else None, _local(sub.submitted_at) if sub else None,
                       c["score"], c["graded_version"], c["coverage"], c["clean"], c["similarity"],
                       " · ".join(c["warnings"])])
            if c["warnings"]:
                for col in range(1, len(dh) + 1):
                    wd.cell(row=wd.max_row, column=col).fill = WARN_FILL
    for row in wd.iter_rows(min_row=2):
        for idx in (4, 7):  # tarih sutunlari
            row[idx].number_format = "dd.mm.yyyy hh:mm"
    _style(wd, len(dh), widths=[14, 22, 22, 30, 17, 16, 10, 17, 8, 10, 12, 11, 14, 50])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _style(ws, ncols: int, widths: list[int]) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = widths[col - 1]
    ws.freeze_panes = "C2"  # baslik ve isim sutunlari sabit


@router.post("/classes/{class_id}/gradebook-link")
def gradebook_link(
    class_id: uuid.UUID,
    db: Session = Depends(get_db),
    user: User = Depends(require_roles("academician", "admin")),
) -> dict:
    cls = _owned_class(db, class_id, user)
    token = create_download_token(str(cls.id), str(user.id), purpose="gradebook")
    return {"url": f"/classes/{cls.id}/gradebook.xlsx?token={token}",
            "expires_in": DOWNLOAD_TTL_MINUTES * 60}


@router.get("/classes/{class_id}/gradebook.xlsx")
def gradebook_download(class_id: uuid.UUID, token: str, db: Session = Depends(get_db)) -> Response:
    try:
        payload = decode_token(token)
    except Exception:
        raise HTTPException(status_code=401, detail="İndirme linkinin süresi dolmuş; yeniden dene.")
    if payload.get("purpose") != "gradebook" or payload.get("sub") != str(class_id):
        raise HTTPException(status_code=401, detail="Geçersiz indirme linki.")
    try:
        user = db.get(User, uuid.UUID(str(payload.get("uid"))))
    except (ValueError, TypeError):
        user = None
    if user is None or user.role not in ("academician", "admin"):
        raise HTTPException(status_code=401, detail="Geçersiz indirme linki.")
    cls = _owned_class(db, class_id, user)  # yetki link acildiginda tekrar dogrulanir
    data = build_gradebook(db, cls)
    filename = f"{_slug(cls.name)}_notlar_{datetime.now(TR):%Y%m%d}.xlsx"
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )

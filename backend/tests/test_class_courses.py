"""Sınıf = bölüm grubu: sınıfa istediği zaman ders eklenir, ödev bir derse ait verilir."""
import io
from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from .conftest import STUDENT_1


def _due(days=3):
    return (datetime.now(timezone.utc) + timedelta(days=days)).isoformat()


def test_class_with_courses_flow(client, aca, s1):
    from .test_logic_review import _extra_user
    t = _extra_user("academician")  # gerçek hoca (demo yeni ders açamaz)
    dep = client.get("/departments", headers=t).json()[0]
    r = client.post("/classes", headers=t, json={"department_id": dep["id"], "name": "Bilgisayar Programcılığı 1"})
    assert r.status_code == 201, r.text
    cls = r.json()
    cid = cls["id"]
    try:
        assert cls["department_name"] == dep["name"] and cls["courses"] == []
        client.post(f"/classes/{cid}/enroll", headers=t, json={"student_no": STUDENT_1})

        # Ders yokken ödev verilemez
        r = client.post("/assignments", headers=t, json={"class_id": cid, "title": "X", "deadline_at": _due()})
        assert r.status_code == 400 and "ders" in r.json()["detail"]

        # Yeni ders adıyla eklenir (bölümde açılır); tek ders varsa ödev ona gider
        r = client.post(f"/classes/{cid}/courses", headers=t, json={"name": "Mesleki Çözümleme I"})
        assert r.status_code == 201, r.text
        mc = next(c for c in r.json()["courses"] if c["name"] == "Mesleki Çözümleme I")
        a1 = client.post("/assignments", headers=t, json={"class_id": cid, "title": "Rapor", "deadline_at": _due()}).json()
        assert a1["course_id"] == mc["id"] and a1["course_name"] == "Mesleki Çözümleme I"

        # Aynı adı tekrar eklemek ikinci ders açmaz
        again = client.post(f"/classes/{cid}/courses", headers=t, json={"name": "mesleki çözümleme i"}).json()
        assert len(again["courses"]) == 1

        # İkinci ders (listeden): artık ders seçmek zorunlu
        other = next(c for c in client.get("/courses", headers=t).json() if c["id"] != mc["id"])
        assert len(client.post(f"/classes/{cid}/courses", headers=t, json={"course_id": other["id"]}).json()["courses"]) == 2
        r = client.post("/assignments", headers=t, json={"class_id": cid, "title": "Y", "deadline_at": _due()})
        assert r.status_code == 400 and "derse" in r.json()["detail"]
        a2 = client.post("/assignments", headers=t, json={"class_id": cid, "title": "Web ödevi",
                                                            "course_id": other["id"], "deadline_at": _due(5)}).json()
        assert a2["course_name"] == other["name"]

        # Sınıfta olmayan ders reddedilir; ödev başka derse taşınabilir
        stray = client.post("/courses", headers=t, json={"department_id": dep["id"], "name": "Başka Ders"}).json()
        r = client.post("/assignments", headers=t, json={"class_id": cid, "title": "Z", "course_id": stray["id"], "deadline_at": _due()})
        assert r.status_code == 400
        moved = client.patch(f"/assignments/{a2['id']}", headers=t, json={"course_id": mc["id"]}).json()
        assert moved["course_name"] == "Mesleki Çözümleme I"
        client.patch(f"/assignments/{a2['id']}", headers=t, json={"course_id": other["id"]})

        # Ödevi olan ders çıkarılamaz; ödevsiz ders çıkarılır
        assert client.delete(f"/classes/{cid}/courses/{mc['id']}", headers=t).status_code == 400
        client.post(f"/classes/{cid}/courses", headers=t, json={"course_id": stray["id"]})
        left = client.delete(f"/classes/{cid}/courses/{stray['id']}", headers=t).json()
        assert {c["id"] for c in left["courses"]} == {mc["id"], other["id"]}

        # Demo hocası yeni ders açamaz (listeyi kirletmesin)
        demo_cls = client.post("/classes", headers=aca, json={"department_id": dep["id"], "name": "Demo"}).json()
        assert client.post(f"/classes/{demo_cls['id']}/courses", headers=aca, json={"name": "Uydurma Ders"}).status_code == 403
        assert client.post(f"/classes/{demo_cls['id']}/courses", headers=aca, json={"name": "Mesleki Çözümleme I"}).status_code == 201
        client.delete(f"/classes/{demo_cls['id']}", headers=aca)

        # Öğrenci ders ekleyemez; sınıf adı değiştirilebilir
        assert client.post(f"/classes/{cid}/courses", headers=s1, json={"name": "Q"}).status_code == 403
        assert client.patch(f"/classes/{cid}", headers=t, json={"name": "Bilgisayar Prog. 1-A"}).json()["name"] == "Bilgisayar Prog. 1-A"

        # Panolar ders adını taşır
        ov = client.get("/me/teaching-overview", headers=t).json()
        c_ov = next(c for c in ov["classes"] if c["id"] == cid)
        assert c_ov["department_name"] == dep["name"] and len(c_ov["courses"]) == 2
        assert {a["course_name"] for a in c_ov["assignments"]} == {"Mesleki Çözümleme I", other["name"]}
        mine = [a for a in client.get("/me/assignments", headers=s1).json() if a["class_id"] == cid]
        assert {a["course_name"] for a in mine} == {"Mesleki Çözümleme I", other["name"]}
        grades = client.get(f"/classes/{cid}/grades", headers=t).json()
        assert all(a["course_name"] for a in grades["assignments"])

        # Excel: her ders ayrı not sayfası
        link = client.post(f"/classes/{cid}/gradebook-link", headers=t).json()["url"]
        wb = load_workbook(io.BytesIO(client.get(link).content))
        assert "Mesleki Çözümleme I" in wb.sheetnames and other["name"][:31] in wb.sheetnames
    finally:
        for a in client.get(f"/assignments?class_id={cid}", headers=t).json():
            client.delete(f"/assignments/{a['id']}", headers=t)
        client.delete(f"/classes/{cid}", headers=t)


def test_legacy_class_create_with_single_course(client, aca):
    """Eski istemci (course_id ile) hâlâ çalışır: bölüm dersten gelir, ders sınıfa eklenir."""
    course = client.get("/courses", headers=aca).json()[0]
    cls = client.post("/classes", headers=aca, json={"course_id": course["id"], "name": "Eski Usul"}).json()
    try:
        assert cls["department_id"] == course["department_id"]
        assert [c["id"] for c in cls["courses"]] == [course["id"]]
    finally:
        client.delete(f"/classes/{cls['id']}", headers=aca)


def test_class_needs_department(client, aca):
    r = client.post("/classes", headers=aca, json={"name": "Bölümsüz"})
    assert r.status_code == 400


def test_teacher_can_add_department_once(client, aca, s1):
    from .test_logic_review import _extra_user
    # Herkese açık demo hocası bölüm listesini kirletemez
    assert client.post("/departments", headers=aca, json={"name": "Deneme"}).status_code == 403
    teacher = _extra_user("academician")
    r = client.post("/departments", headers=teacher, json={"name": "Elektrik  Teknolojisi"})
    assert r.status_code == 201 and r.json()["name"] == "Elektrik Teknolojisi"
    again = client.post("/departments", headers=teacher, json={"name": "elektrik teknolojisi"}).json()
    assert again["id"] == r.json()["id"]
    assert client.post("/departments", headers=s1, json={"name": "X"}).status_code == 403

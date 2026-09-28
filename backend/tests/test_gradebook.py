"""Sinif not dosyasi (Excel): icerik, bos hucreler, uyarilar, indirme linki guvenligi."""
import io
from datetime import datetime, timedelta, timezone

import pytest
from openpyxl import load_workbook

from .conftest import STUDENT_1, STUDENT_2


@pytest.fixture
def gb_class(client, aca):
    course = client.get("/courses", headers=aca).json()[0]
    cls = client.post("/classes", headers=aca, json={"course_id": course["id"], "name": "Not Sınıfı"}).json()
    for no in (STUDENT_1, STUDENT_2):
        client.post(f"/classes/{cls['id']}/enroll", headers=aca, json={"student_no": no})
    return cls["id"]


def _assignment(client, aca, cid, title, days):
    return client.post("/assignments", headers=aca, json={
        "class_id": cid, "title": title, "requirements": ["A olmalı"],
        "deadline_at": (datetime.now(timezone.utc) + timedelta(days=days)).isoformat(),
    }).json()["id"]


def _download(client, headers, cid):
    link = client.post(f"/classes/{cid}/gradebook-link", headers=headers)
    assert link.status_code == 200, link.text
    r = client.get(link.json()["url"])
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    return load_workbook(io.BytesIO(r.content))


def test_gradebook_contents(client, aca, s1, s2, gb_class, submit):
    a1 = _assignment(client, aca, gb_class, "Ödev 1", 1)
    a2 = _assignment(client, aca, gb_class, "Ödev 2", 2)
    v1 = submit(s1, a1, {"main.py": "x = 1\n"})
    client.post(f"/submissions/{v1['id']}/score", headers=aca, json={"score": 80})
    submit(s1, a1, {"main.py": "x = 2\n"})  # notlandiktan sonra yeni surum
    b = submit(s1, a2, {"main.py": "y = 1\n"})
    client.post(f"/submissions/{b['id']}/score", headers=aca, json={"score": 90})
    submit(s2, a2, {"main.py": "z = 1\n"})  # notlanmadi; s2 Odev 1'i teslim etmedi

    wb = _download(client, aca, gb_class)
    notes = wb["Notlar"]
    rows = list(notes.iter_rows(values_only=True))
    assert rows[0] == ("Öğrenci No", "Ad Soyad", "Ödev 1", "Ödev 2", "Ortalama*")
    assert rows[1] == (STUDENT_1, "Mehmet Demir", 80, 90, 85.0)
    assert rows[2] == (STUDENT_2, "Zeynep Kaya", None, None, None)  # 0 degil, bos
    assert "Ortalama yalnızca" in rows[4][0]

    # Ayrıntı: Ders sütunu ödevden önce
    detail = {(r[0], r[3]): r for r in wb["Ayrıntı"].iter_rows(min_row=2, values_only=True)}
    d = detail[(STUDENT_1, "Ödev 1")]
    assert d[2]  # dersin adı
    assert d[5] == "Zamanında" and d[6] == 2 and d[8] == 80 and d[9] == 1
    assert "Notlandıktan sonra yeni sürüm" in d[13]
    assert isinstance(d[7], datetime)
    assert detail[(STUDENT_2, "Ödev 1")][5] == "Teslim yok"
    assert detail[(STUDENT_2, "Ödev 2")][8] is None


def test_gradebook_link_security(client, aca, s1, gb_class, submit, class_id):
    assert client.post(f"/classes/{gb_class}/gradebook-link", headers=s1).status_code == 403
    url = client.post(f"/classes/{gb_class}/gradebook-link", headers=aca).json()["url"]
    token = url.split("token=")[1]
    # Baska sinifin dosyasini acmaz; oturum token'i yerine gecmez
    assert client.get(f"/classes/{class_id}/gradebook.xlsx?token={token}").status_code == 401
    session = aca["Authorization"].split(" ")[1]
    assert client.get(f"/classes/{gb_class}/gradebook.xlsx?token={session}").status_code == 401
    assert client.get("/classes", headers={"Authorization": f"Bearer {token}"}).status_code == 401
    # Proje indirme linki not dosyasini acmaz (amac farkli)
    aid = _assignment(client, aca, gb_class, "Ödev X", 3)
    sub = submit(s1, aid, {"main.py": "x = 1\n"})
    dl = client.post(f"/submissions/{sub['id']}/download-link", headers=aca).json()["url"]
    dl_token = dl.split("token=")[1]
    assert client.get(f"/classes/{gb_class}/gradebook.xlsx?token={dl_token}").status_code == 401


def test_empty_class_still_produces_file(client, aca, gb_class):
    wb = _download(client, aca, gb_class)
    rows = list(wb["Notlar"].iter_rows(values_only=True))
    assert rows[0] == ("Öğrenci No", "Ad Soyad", "Ortalama*")
    assert len(rows) >= 3  # iki ogrenci satiri
